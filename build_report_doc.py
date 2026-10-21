import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DOCX_PATH = os.path.join(BASE_DIR, "CampusSafe_BTech_Major_Project_Report.docx")
LOGO_PATH = os.path.join(BASE_DIR, "bvp_logo.jpg")
FIG_2_1 = os.path.join(BASE_DIR, "figure_2_1.png")
FIG_4_1 = os.path.join(BASE_DIR, "figure_4_1.png")
FIG_6_1 = os.path.join(BASE_DIR, "arch_diagram.png")
FIG_6_2 = os.path.join(BASE_DIR, "fsm_erd_diagram.png")
FIG_9_1 = os.path.join(BASE_DIR, "results_chart.png")

# Color constants
C_NAVY = RGBColor(0x0F, 0x2A, 0x4A)
C_DARK = RGBColor(0x1E, 0x29, 0x3B)
C_MUTED = RGBColor(0x47, 0x55, 0x69)

def create_document():
    doc = docx.Document()
    
    # Configure default style font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = C_DARK
    
    # Configure initial section margins: Top 1", Bottom 1", Left 1.25", Right 1"
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.0)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    
    return doc

def add_p(doc, text="", font_size=12, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, 
          space_before=0, space_after=6, line_spacing=1.5, indent=False, color=C_DARK):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.3)
        
    if text:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
    return p

def add_chapter_title(doc, chapter_num, title_text):
    p_num = add_p(doc, f"CHAPTER {chapter_num}", font_size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=6)
    p_title = add_p(doc, title_text.upper(), font_size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=18)
    return p_title

def add_section_heading(doc, num_str, title_text):
    text = f"{num_str}  {title_text.upper()}"
    return add_p(doc, text, font_size=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=14, space_after=6)

def add_subsection_heading(doc, num_str, title_text):
    text = f"{num_str}  {title_text}"
    return add_p(doc, text, font_size=12, bold=True, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=10, space_after=4)

def add_bullet_item(doc, lead_in, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    
    r_lead = p.add_run(lead_in + ": ")
    r_lead.font.name = 'Times New Roman'
    r_lead.font.size = Pt(12)
    r_lead.font.bold = True
    r_lead.font.color.rgb = C_DARK
    
    r_body = p.add_run(text)
    r_body.font.name = 'Times New Roman'
    r_body.font.size = Pt(12)
    r_body.font.bold = False
    r_body.font.color.rgb = C_DARK
    return p

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def format_table(table, header_bg='1E3A8A'):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Format header row
    hdr = table.rows[0]
    for cell in hdr.cells:
        set_cell_background(cell, header_bg)
        set_cell_margins(cell, 120, 120, 150, 150)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                
    # Format data rows
    for r_idx, row in enumerate(table.rows[1:]):
        bg = 'F8FAFC' if r_idx % 2 == 1 else 'FFFFFF'
        for cell in row.cells:
            if bg != 'FFFFFF':
                set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 120, 120)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = C_DARK

def add_figure(doc, img_path, caption, width_in=5.8):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(12)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(img_path, width=Inches(width_in))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(14)
        run_cap = p_cap.add_run(caption)
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(10.5)
        run_cap.font.bold = True
        run_cap.font.color.rgb = C_MUTED
    else:
        print(f"Warning: Figure not found at {img_path}")

print("Framework helpers initialized successfully.")
